"""ไฟล์ Excel งบประมาณเบื้องต้นจาก BOM

ตั้งใจให้เปิดแล้วเห็นทั้งตัวเลขและ "ความน่าเชื่อถือของตัวเลข" ในหน้าเดียว —
รายการที่ยังไม่มีราคา หรือใช้ราคาเก่าเกินหนึ่งปี ถูกทำสีไว้ให้เห็นชัด
ไม่ใช่ซ่อนไว้ให้ยอดรวมดูสวย
"""
import io
from datetime import datetime
from typing import Any, Dict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.config import settings

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _as_bytes(wb: Workbook, filename: str) -> Dict[str, Any]:
    """คืนไฟล์เป็น bytes ไม่เขียนลงดิสก์ — เก็บลง MongoDB ผ่าน services/filestore.py"""
    buffer = io.BytesIO()
    wb.save(buffer)
    return {"filename": filename, "data": buffer.getvalue(), "content_type": XLSX_TYPE}

THIN = Side(style="thin", color="B7C0CD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3B63")
WARN_FILL = PatternFill("solid", fgColor="FBE9E7")     # ยังไม่มีราคา
STALE_FILL = PatternFill("solid", fgColor="F6ECCD")    # ราคาเก่า
TOTAL_FILL = PatternFill("solid", fgColor="EEF2F7")

TITLE_FONT = Font(size=15, bold=True, color="1F3B63")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
SMALL = Font(size=9, color="6B7280")
BOLD = Font(bold=True)

MONEY = "#,##0.00"

CONF_LABEL = {"high": "ตรงมาก", "medium": "น่าจะใช่", "low": "ไม่แน่ใจ"}
SOURCE_LABEL = {
    "last_price": "ราคาซื้อล่าสุด",
    "manual": "กรอกเอง",
    "none": "ยังไม่มีราคา",
}


def _fmt_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    text = str(value or "")
    return text[:10] if text else "-"


def build_bom_workbook(bom: Dict[str, Any]) -> Dict[str, Any]:
    totals = bom.get("totals") or {}
    lines = bom.get("lines") or []

    wb = Workbook()
    ws = wb.active
    ws.title = "งบประมาณเบื้องต้น"

    columns = [
        ("ลำดับ", 7), ("รายการตาม BOM", 40), ("จำนวน", 10), ("หน่วย", 9),
        ("รหัสสินค้าที่จับคู่", 20), ("ชื่อสินค้าในระบบ", 40), ("ความมั่นใจ", 12),
        ("ราคา/หน่วย", 14), ("ที่มาของราคา", 16), ("ราคา ณ วันที่", 14),
        ("อายุราคา (วัน)", 14), ("จำนวนเงิน", 16), ("หมายเหตุ", 26),
    ]
    for idx, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    last_col = len(columns)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(row=1, column=1, value="งบประมาณเบื้องต้นจาก BOM").font = TITLE_FONT
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(row=2, column=1, value="{} · {} · ออกเมื่อ {}".format(
        settings.APP_NAME, bom.get("bom_no", ""), _fmt_date(datetime.now())
    )).font = SMALL
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    ws.cell(row=4, column=1, value="โครงการ").font = BOLD
    ws.cell(row=4, column=2, value=bom.get("title", ""))
    ws.cell(row=5, column=1, value="หมายเหตุ").font = BOLD
    ws.cell(row=5, column=2, value=bom.get("note", "") or "-")

    # เตือนเรื่องความครอบคลุมไว้บนสุด ก่อนที่ใครจะเลื่อนไปดูยอดรวม
    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=last_col)
    warn = ws.cell(row=6, column=1, value=(
        "ตีราคาได้ {}/{} รายการ ({}%) · ยังไม่มีราคา {} รายการ · "
        "ใช้ราคาเก่ากว่า 1 ปี {} รายการ · จับคู่แบบไม่แน่ใจ {} รายการ — "
        "ยอดรวมด้านล่างคิดเฉพาะรายการที่มีราคาเท่านั้น"
    ).format(
        totals.get("priced_lines", 0), totals.get("line_count", 0),
        totals.get("coverage_percent", 0), totals.get("unpriced_lines", 0),
        totals.get("stale_price_lines", 0), totals.get("low_confidence_lines", 0),
    ))
    warn.font = Font(size=10, bold=True, color="8A5A00")
    warn.fill = STALE_FILL
    warn.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[6].height = 30

    header_row = 8
    for idx, (label, _) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=idx, value=label)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row = header_row + 1
    for line in lines:
        values = [
            line.get("line_no"),
            line.get("name", ""),
            line.get("qty"),
            line.get("uom", ""),
            line.get("part_num") or "-",
            line.get("item_description") or "-",
            CONF_LABEL.get(line.get("confidence"), "-") if line.get("matched") else "ยังไม่จับคู่",
            line.get("unit_price"),
            SOURCE_LABEL.get(line.get("price_source"), "-"),
            _fmt_date(line.get("price_date")),
            line.get("price_age_days"),
            line.get("amount"),
            line.get("remark", ""),
        ]
        for idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=idx, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top", wrap_text=idx in (2, 6, 13),
                horizontal="right" if idx in (3, 8, 11, 12) else None,
            )
            if idx in (8, 12):
                cell.number_format = MONEY
        if not line.get("priced"):
            for idx in range(1, last_col + 1):
                ws.cell(row=row, column=idx).fill = WARN_FILL
        elif line.get("price_stale"):
            for idx in range(1, last_col + 1):
                ws.cell(row=row, column=idx).fill = STALE_FILL
        row += 1

    def total_row(label: str, value: Any, bold: bool = False) -> None:
        nonlocal row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col - 2)
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.alignment = Alignment(horizontal="right")
        label_cell.fill = TOTAL_FILL
        if bold:
            label_cell.font = BOLD
        value_cell = ws.cell(row=row, column=last_col - 1, value=value)
        value_cell.number_format = MONEY
        value_cell.fill = TOTAL_FILL
        value_cell.border = BORDER
        if bold:
            value_cell.font = BOLD
        row += 1

    row += 1
    total_row("รวมค่าของตามราคาซื้อล่าสุด", totals.get("subtotal"))
    total_row("เผื่อสำรอง {}%".format(totals.get("contingency_percent", 0)),
              totals.get("contingency_amount"))
    total_row("รวมก่อนภาษี", totals.get("before_vat"))
    total_row("ภาษีมูลค่าเพิ่ม {}%".format(totals.get("vat_percent", 0)),
              totals.get("vat_amount"))
    total_row("รวมงบประมาณทั้งสิ้น ({})".format(bom.get("currency", "THB")),
              totals.get("total"), bold=True)

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.cell(row=row, column=1, value=(
        "ตัวเลขนี้เป็นการประมาณจากราคาที่บริษัทเคยซื้อจริง ไม่ใช่ราคาเสนอจากผู้ขาย "
        "ใช้เพื่อตั้งงบเบื้องต้นเท่านั้น ราคาจริงต้องขอใบเสนอราคาอีกครั้ง"
    )).font = SMALL

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_title_rows = "{}:{}".format(header_row, header_row)

    safe = "".join(ch for ch in (bom.get("bom_no") or "BOM") if ch.isalnum() or ch in "-_")
    return _as_bytes(wb, "{}_budget.xlsx".format(safe))
