"""สร้างเอกสาร Excel: ใบขอราคา (RFQ) ที่ส่งแนบให้ผู้ขาย และตารางเปรียบเทียบราคา

ใบขอราคาหนึ่งไฟล์ต่อผู้ขายหนึ่งราย ช่องสีเหลืองคือช่องที่ผู้ขายต้องกรอก
คอลัมน์ "จำนวนเงิน" ใส่สูตร Excel ไว้ให้คิดเองเมื่อกรอกราคาต่อหน่วย
"""
import io
from datetime import date, datetime
from typing import Any, Dict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.config import settings

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _as_bytes(wb: Workbook, filename: str) -> Dict[str, Any]:
    """คืนไฟล์เป็น bytes ไม่เขียนลงดิสก์ — คนเรียกเป็นคนตัดสินใจว่าจะเก็บที่ไหน
    (ตอนนี้เก็บลง MongoDB ผ่าน services/filestore.py)"""
    buffer = io.BytesIO()
    wb.save(buffer)
    return {"filename": filename, "data": buffer.getvalue(), "content_type": XLSX_TYPE}

THIN = Side(style="thin", color="B7C0CD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3B63")
FILL_ME = PatternFill("solid", fgColor="FFF7DC")     # ช่องที่ผู้ขายต้องกรอก
LABEL_FILL = PatternFill("solid", fgColor="EEF2F7")

TITLE_FONT = Font(size=16, bold=True, color="1F3B63")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
LABEL_FONT = Font(bold=True, size=10)
SMALL = Font(size=9, color="6B7280")


def _fmt_date(value) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    return str(value or "-")


def _put_label(ws, row: int, col: int, label: str, value) -> None:
    cell = ws.cell(row=row, column=col, value=label)
    cell.font = LABEL_FONT
    cell.fill = LABEL_FILL
    cell.border = BORDER
    v = ws.cell(row=row, column=col + 1, value=value if value not in (None, "") else "-")
    v.border = BORDER
    v.alignment = Alignment(vertical="center", wrap_text=True)


def build_rfq_workbook(
    rfq: Dict[str, Any], vendor: Dict[str, Any], invite: Dict[str, Any], message: str = ""
) -> Dict[str, Any]:
    """สร้างไฟล์ใบขอราคาสำหรับผู้ขายรายหนึ่ง — คอลัมน์ราคาเว้นว่างให้กรอก"""
    wb = Workbook()
    ws = wb.active
    ws.title = "RFQ"

    columns = [
        ("ลำดับ", 7), ("รหัสสินค้า", 20), ("รายละเอียด", 42), ("ผู้ผลิต / MPN", 24),
        ("สเปก", 30), ("จำนวน", 10), ("หน่วย", 9),
        ("ราคาต่อหน่วย *", 15), ("จำนวนเงิน", 15), ("MOQ", 10),
        ("Lead time (วัน)", 14), ("รหัสของผู้ขาย", 18), ("หมายเหตุ", 26),
    ]
    for idx, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    last_col = len(columns)

    # ---------- หัวเอกสาร ----------
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(row=1, column=1, value="ใบขอเสนอราคา / REQUEST FOR QUOTATION").font = TITLE_FONT
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(row=2, column=1, value=settings.APP_NAME).font = SMALL
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    row = 4
    _put_label(ws, row, 1, "เลขที่ RFQ", rfq.get("rfq_no", ""))
    _put_label(ws, row, 4, "วันที่ออกเอกสาร", _fmt_date(datetime.now()))
    _put_label(ws, row, 7, "สกุลเงิน", rfq.get("currency", "THB"))
    row += 1
    _put_label(ws, row, 1, "เรื่อง", rfq.get("title", ""))
    _put_label(ws, row, 4, "กำหนดยื่นราคา", _fmt_date(rfq.get("due_date")))
    _put_label(ws, row, 7, "Incoterm", rfq.get("incoterm", ""))
    row += 1
    _put_label(ws, row, 1, "เรียน (ผู้ขาย)", vendor.get("name", ""))
    _put_label(ws, row, 4, "ผู้ติดต่อ", invite.get("contact_name") or invite.get("contact_email", ""))
    _put_label(ws, row, 7, "เงื่อนไขชำระเงิน", rfq.get("payment_terms", ""))
    row += 1
    _put_label(ws, row, 1, "สถานที่ส่งมอบ", rfq.get("delivery_place", ""))
    _put_label(ws, row, 4, "อ้างอิง BOM", rfq.get("bom_code", ""))
    _put_label(ws, row, 7, "จำนวนรายการ", len(rfq.get("lines", [])))

    row += 2
    note = message or rfq.get("note") or "กรุณากรอกราคาในช่องสีเหลือง แล้วส่งกลับพร้อมใบเสนอราคาของบริษัทท่าน"
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=f"หมายเหตุ: {note}")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 30

    # ---------- ตารางรายการ ----------
    row += 2
    header_row = row
    for idx, (label, _) in enumerate(columns, start=1):
        c = ws.cell(row=header_row, column=idx, value=label)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30

    for i, line in enumerate(rfq.get("lines", []), start=1):
        r = header_row + i
        specs = line.get("specs") or {}
        spec_text = ", ".join(f"{k}: {v}" for k, v in specs.items())
        mfr = " / ".join(x for x in (line.get("manufacturer", ""), line.get("mpn", "")) if x)

        values = [
            i, line.get("item_code", ""), line.get("description") or line.get("name", ""),
            mfr, spec_text, line.get("qty", 0), line.get("uom", ""),
            None, None, None, None, None, line.get("remark", ""),
        ]
        for col_idx, value in enumerate(values, start=1):
            c = ws.cell(row=r, column=col_idx, value=value)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=col_idx in (3, 5, 13))
            if col_idx in (8, 10, 11, 12):        # ช่องที่ผู้ขายกรอก
                c.fill = FILL_ME
            if col_idx in (6, 8, 9):
                c.number_format = "#,##0.00"
        # จำนวนเงิน = จำนวน x ราคาต่อหน่วย (คิดให้อัตโนมัติเมื่อกรอกราคา)
        ws.cell(row=r, column=9).value = f"=IF(H{r}=\"\",\"\",F{r}*H{r})"
        ws.cell(row=r, column=9).fill = PatternFill("solid", fgColor="F3F6FA")

    n_lines = len(rfq.get("lines", []))
    total_row = header_row + n_lines + 1
    ws.cell(row=total_row, column=8, value="รวมเป็นเงิน").font = LABEL_FONT
    ws.cell(row=total_row, column=8).alignment = Alignment(horizontal="right")
    total = ws.cell(
        row=total_row, column=9,
        value=f"=SUM(I{header_row + 1}:I{header_row + n_lines})" if n_lines else 0,
    )
    total.font = LABEL_FONT
    total.number_format = "#,##0.00"
    total.border = BORDER

    # ---------- ช่องข้อมูลผู้เสนอราคา ----------
    row = total_row + 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.cell(row=row, column=1, value="ส่วนของผู้เสนอราคา (กรุณากรอก)").font = Font(bold=True, size=11, color="1F3B63")

    for label in (
        ("ชื่อผู้ติดต่อ", "อีเมล", "โทรศัพท์"),
        ("ยืนราคาถึงวันที่", "ระยะเวลาส่งมอบ (วัน)", "เงื่อนไขชำระเงิน"),
    ):
        row += 1
        for i, text in enumerate(label):
            col = 1 + i * 3
            c = ws.cell(row=row, column=col, value=text)
            c.font = LABEL_FONT
            c.fill = LABEL_FILL
            c.border = BORDER
            blank = ws.cell(row=row, column=col + 1)
            blank.fill = FILL_ME
            blank.border = BORDER

    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    portal = invite.get("portal_url", "")
    ws.cell(
        row=row, column=1,
        value=f"กรอกราคาออนไลน์และแนบใบเสนอราคาได้ที่: {portal}" if portal else "",
    ).font = Font(size=10, color="1F3B63", underline="single")

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_title_rows = f"{header_row}:{header_row}"

    safe_vendor = "".join(
        ch for ch in (vendor.get("vendor_key") or "vendor").replace(":", "-") if ch.isalnum() or ch in "-_"
    )
    return _as_bytes(wb, f"{rfq.get('rfq_no', 'RFQ')}_{safe_vendor}.xlsx")


def build_comparison_workbook(comparison: Dict[str, Any]) -> Dict[str, Any]:
    """ตารางเปรียบเทียบราคาจากผู้ขายทุกราย — ราคาต่ำสุดของแต่ละแถวถูกไฮไลต์"""
    wb = Workbook()
    ws = wb.active
    ws.title = "เปรียบเทียบราคา"

    vendors = comparison.get("vendors", [])
    rows = comparison.get("rows", [])

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5 + len(vendors) * 2)
    ws.cell(
        row=1, column=1,
        value=f"ตารางเปรียบเทียบราคา {comparison.get('rfq_no', '')} — {comparison.get('title', '')}",
    ).font = TITLE_FONT

    header = ["ลำดับ", "รหัสสินค้า", "รายละเอียด", "จำนวน", "หน่วย"]
    for v in vendors:
        label = (v.get("vendor_name") or v["vendor_key"])[:24]
        header += [f"{label}\nราคา/หน่วย", f"{label}\nรวม"]
    header.append("ราคาต่ำสุดของ")

    head_row = 3
    for idx, label in enumerate(header, start=1):
        c = ws.cell(row=head_row, column=idx, value=label)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["B"].width = 20

    best_fill = PatternFill("solid", fgColor="D6F5DE")
    for i, row_data in enumerate(rows, start=1):
        r = head_row + i
        base = [i, row_data["item_code"], row_data["name"], row_data["qty"], row_data["uom"]]
        for col_idx, value in enumerate(base, start=1):
            c = ws.cell(row=r, column=col_idx, value=value)
            c.border = BORDER

        cells_by_vendor = {c["vendor_key"]: c for c in row_data.get("cells", [])}
        for vi, v in enumerate(vendors):
            cell_data = cells_by_vendor.get(v["vendor_key"], {})
            price_col = 6 + vi * 2
            unit = ws.cell(row=r, column=price_col, value=cell_data.get("unit_price"))
            amount = ws.cell(row=r, column=price_col + 1, value=cell_data.get("amount"))
            for c in (unit, amount):
                c.border = BORDER
                c.number_format = "#,##0.00"
                if cell_data.get("is_lowest"):
                    c.fill = best_fill
        winner = next(
            (v.get("vendor_name", "") for v in vendors if v["vendor_key"] == row_data.get("lowest_vendor_key")),
            "",
        )
        ws.cell(row=r, column=6 + len(vendors) * 2, value=winner).border = BORDER

    total_row = head_row + len(rows) + 1
    ws.cell(row=total_row, column=5, value="รวมทั้งสิ้น").font = LABEL_FONT
    for vi, v in enumerate(vendors):
        c = ws.cell(row=total_row, column=7 + vi * 2, value=v.get("total"))
        c.font = LABEL_FONT
        c.number_format = "#,##0.00"
        c.border = BORDER
        if v["vendor_key"] == comparison.get("best_total_vendor_key"):
            c.fill = best_fill

    ws.freeze_panes = ws.cell(row=head_row + 1, column=6)

    return _as_bytes(wb, f"{comparison.get('rfq_no', 'RFQ')}_comparison.xlsx")
